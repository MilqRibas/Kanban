"""Parse real Suprema TX XLSX samples and emit recovery JSON for CRM incentives."""
from __future__ import annotations

import json
import re
from collections import Counter
from datetime import datetime, date, time, timezone
from pathlib import Path

from openpyxl import load_workbook

SAMPLES = Path(r"c:\Users\sxgru\OneDrive\Documentos\Projetos\Kanban\docs\samples")
INPUT = SAMPLES / "_fixed"
if not INPUT.exists():
    INPUT = SAMPLES
MKT = "1092502"


def norm_header(v) -> str:
    s = str(v or "").strip().lower()
    s = (
        s.replace("á", "a")
        .replace("à", "a")
        .replace("ã", "a")
        .replace("â", "a")
        .replace("é", "e")
        .replace("ê", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ô", "o")
        .replace("õ", "o")
        .replace("ú", "u")
        .replace("ç", "c")
    )
    return re.sub(r"\s+", " ", s)


FIELD_MAP = {
    "id": "external_transaction_id",
    "transaction id": "external_transaction_id",
    "receiver player id": "receiver_player_id",
    "sender player id": "sender_player_id",
    "sender player nickname": "sender_nickname",
    "receiver nickname": "receiver_nickname",
    "receiver player nickname": "receiver_nickname",
    "agente player id": "agent_id",
    "agent player id": "agent_id",
    "agent id": "agent_id",
    "dia": "day",
    "hora": "time",
    "origem": "origin",
    "sx tipo": "sx_type",
    "chips send out": "chips_send_out",
    "chips claimback": "chips_claimback",
    "status sistema": "system_status",
    "system status": "system_status",
    "order status": "order_status",
}


def cell_id(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, float):
        if not (v == v):  # NaN
            return None
        if v.is_integer():
            return str(int(v))
        return str(int(v)) if abs(v - round(v)) < 1e-9 else str(v).rstrip("0").rstrip(".")
    if isinstance(v, int):
        return str(v)
    s = str(v).strip()
    if not s or s == "-":
        return None
    try:
        f = float(s.replace(",", "."))
        if abs(f - round(f)) < 1e-9:
            return str(int(round(f)))
    except Exception:
        pass
    return s


def cell_num(v):
    if v is None or v == "" or v == "-":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except Exception:
        return None


def cell_text(v):
    if v is None:
        return None
    s = str(v).strip()
    return None if (not s or s == "-") else s


def combine_day_time(day_v, time_v):
    d = None
    t = None
    if isinstance(day_v, datetime):
        d = day_v.date()
        if day_v.hour or day_v.minute or day_v.second:
            t = day_v.time()
    elif isinstance(day_v, date):
        d = day_v
    elif day_v is not None:
        s = str(day_v).strip()
        m = re.match(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", s)
        if m:
            y = m.group(3)
            if len(y) == 2:
                y = "20" + y
            d = date(int(y), int(m.group(2)), int(m.group(1)))
    if isinstance(time_v, datetime):
        t = time_v.time()
    elif isinstance(time_v, time):
        t = time_v
    elif time_v is not None:
        s = str(time_v).strip()
        m = re.match(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", s)
        if m:
            t = time(int(m.group(1)), int(m.group(2)), int(m.group(3) or 0))
    if not d:
        return None
    if not t:
        t = time(0, 0, 0)
    return datetime(d.year, d.month, d.day, t.hour, t.minute, t.second, tzinfo=timezone.utc).isoformat().replace(
        "+00:00", "Z"
    )


def is_bonus(sx_type: str | None) -> bool:
    if not sx_type:
        return False
    n = norm_header(sx_type).replace("ô", "o")
    return "bonus" in n


def parse_file(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header_row = None
    colmap = {}
    labels = {}
    for i, row in enumerate(rows):
        cells = list(row)
        norms = [norm_header(c) for c in cells]
        mapped = {}
        labs = {}
        for idx, h in enumerate(norms):
            if h in FIELD_MAP and FIELD_MAP[h] not in mapped:
                mapped[FIELD_MAP[h]] = idx
                labs[FIELD_MAP[h]] = str(cells[idx] or "").strip()
        if "external_transaction_id" in mapped and "receiver_player_id" in mapped:
            header_row = i
            colmap = mapped
            labels = labs
            break
    if header_row is None:
        wb.close()
        raise RuntimeError(f"headers not found in {path.name}")

    transactions = []
    for row in rows:
        cells = list(row)
        if not any(c is not None and str(c).strip() not in ("", "-") for c in cells):
            continue

        def get(field):
            idx = colmap.get(field)
            return None if idx is None or idx >= len(cells) else cells[idx]

        ext = cell_id(get("external_transaction_id"))
        recv = cell_id(get("receiver_player_id"))
        if not ext or not recv:
            continue
        sx = cell_text(get("sx_type"))
        chips = cell_num(get("chips_send_out"))
        amount = chips if chips is not None else 0.0
        sender = cell_id(get("sender_player_id"))
        bonus = is_bonus(sx)
        transactions.append(
            {
                "external_transaction_id": ext,
                "receiver_player_id": recv,
                "receiver_nickname": cell_text(get("receiver_nickname")),
                "sender_player_id": sender,
                "sender_nickname": cell_text(get("sender_nickname")),
                "agent_id": cell_id(get("agent_id")),
                "occurred_at": combine_day_time(get("day"), get("time")),
                "origin": cell_text(get("origin")),
                "transaction_type": sx,
                "amount": amount,
                "chips_send_out": chips,
                "chips_claimback": cell_num(get("chips_claimback")),
                "system_status": cell_text(get("system_status")),
                "order_status": cell_text(get("order_status")),
                "is_bonus": bonus,
                "is_deposit": norm_header(cell_text(get("origin")) or "") in ("sx 24 horas", "sx 24 horas"),
            }
        )
        # deposit: origem SX 24 Horas
        origin_n = norm_header(cell_text(get("origin")) or "")
        transactions[-1]["is_deposit"] = origin_n.replace("ô", "o") == "sx 24 horas" and not bonus

    wb.close()
    mkt = [t for t in transactions if t["sender_player_id"] == MKT]
    with_sender = [t for t in transactions if t["sender_player_id"]]
    incentives = [t for t in transactions if t["sender_player_id"] == MKT or t["is_bonus"]]
    return {
        "file": path.name,
        "labels": labels,
        "rows": len(transactions),
        "withSender": len(with_sender),
        "mktGt": len(mkt),
        "mktGtAmount": round(sum(abs(t["amount"]) for t in mkt), 2),
        "incentives": len(incentives),
        "incentiveAmount": round(sum(abs(t["amount"]) for t in incentives), 2),
        "uniqueSenders": sorted({t["sender_player_id"] for t in with_sender}),
        "senderNicknameTop": Counter(
            t["sender_nickname"] for t in mkt if t.get("sender_nickname")
        ).most_common(5),
        "transactions": transactions,
        "mktSample": mkt[:5],
    }


def main():
    files = sorted(INPUT.glob("*.xlsx"))
    reports = []
    all_tx = []
    sender_updates = []
    for path in files:
        print("parsing", path.name)
        rep = parse_file(path)
        reports.append({k: v for k, v in rep.items() if k != "transactions"})
        for t in rep["transactions"]:
            row = dict(t)
            row["source_file"] = path.name
            all_tx.append(row)
            if t["sender_player_id"]:
                sender_updates.append(
                    {
                        "external_transaction_id": t["external_transaction_id"],
                        "sender_player_id": t["sender_player_id"],
                        "sender_nickname": t.get("sender_nickname"),
                    }
                )
        print(
            " ",
            rep["rows"],
            "rows; senderHeader=",
            rep["labels"].get("sender_player_id"),
            "; mkt=",
            rep["mktGt"],
            rep["mktGtAmount"],
        )

    # dedupe sender updates by external id (last wins)
    by_id = {}
    for u in sender_updates:
        by_id[u["external_transaction_id"]] = u
    sender_updates = list(by_id.values())

    # dedupe all_tx by external id preferring rows with sender
    by_tx = {}
    for t in all_tx:
        ext = t["external_transaction_id"]
        prev = by_tx.get(ext)
        if prev is None or (t.get("sender_player_id") and not prev.get("sender_player_id")):
            by_tx[ext] = t
    all_tx = list(by_tx.values())

    summary = {
        "files": len(files),
        "reports": reports,
        "totalRows": len(all_tx),
        "uniqueExtIds": len(all_tx),
        "senderUpdates": len(sender_updates),
        "mktGtUnique": sum(1 for t in all_tx if t.get("sender_player_id") == MKT),
        "mktGtAmountUnique": round(
            sum(abs(t["amount"]) for t in all_tx if t.get("sender_player_id") == MKT), 2
        ),
        "mktGtTotalRawSum": sum(r["mktGt"] for r in reports),
        "mktGtAmountTotalRawSum": round(sum(r["mktGtAmount"] for r in reports), 2),
    }
    (SAMPLES / "parse-all-summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (SAMPLES / "sender-updates.json").write_text(
        json.dumps(sender_updates, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (SAMPLES / "parsed-all-rows.json").write_text(
        json.dumps(all_tx, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps({k: summary[k] for k in summary if k != "reports"}, indent=2))
    for r in reports:
        print(
            f"{r['file']}: rows={r['rows']} withSender={r['withSender']} mkt={r['mktGt']} amt={r['mktGtAmount']} header={r['labels'].get('sender_player_id')}"
        )


if __name__ == "__main__":
    main()
